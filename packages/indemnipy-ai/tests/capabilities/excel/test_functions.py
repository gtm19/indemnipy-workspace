from datetime import datetime, timezone
from typing import Any

import pytest
from indemnipy_ai.capabilities.excel._functions import (
    DateParsingOptions,
    _cleanse_data,
    _dataframe_from_range,
)
from indemnipy_ai.capabilities.excel._functions import (
    _gently_parse_datetime as gently_parse_datetime,
)
from inline_snapshot import snapshot
from openpyxl import Workbook as OpxlWorkbook


@pytest.mark.parametrize(
    "relaxed_about_day",
    [True, False],
    ids=["relaxed_about_day=True", "relaxed_about_day=False"],
)
@pytest.mark.parametrize(
    "input_value, expected_output",
    [
        ("2023-01-15", datetime(2023, 1, 15, tzinfo=timezone.utc)),
        ("15/01/2023", datetime(2023, 1, 15, tzinfo=timezone.utc)),
        ("January 15, 2023", datetime(2023, 1, 15, tzinfo=timezone.utc)),
        ("15 Jan 2023", datetime(2023, 1, 15, tzinfo=timezone.utc)),
        ("01-15-2023", datetime(2023, 1, 15, tzinfo=timezone.utc)),
        ("2023/01/15 4:10:12", datetime(2023, 1, 15, 4, 10, 12, tzinfo=timezone.utc)),
        ("2023/01/15 4:30am", datetime(2023, 1, 15, 4, 30, tzinfo=timezone.utc)),
        ("2023.01.15 4am", datetime(2023, 1, 15, 4, tzinfo=timezone.utc)),
    ],
)
def test_should_always_parse(
    input_value: Any, expected_output: datetime, relaxed_about_day: bool
):
    """
    Tests for inputs which should always parsed as datetimes regardless of the relaxed_about_day setting.
    """
    result = gently_parse_datetime(input_value, relaxed_about_day=relaxed_about_day)
    assert result == expected_output


@pytest.mark.parametrize(
    "relaxed_about_day",
    [True, False],
    ids=["relaxed_about_day=True", "relaxed_about_day=False"],
)
@pytest.mark.parametrize(
    "input_value",
    [
        "2/4/2023",  # Ambiguous date, should return original string
        "2023-13-01",  # Invalid month, should return original string
        "2023-02-30",  # Invalid day, should return original string
        "Back in 2021",  # No month or day, should return original string
        "Back in year of yonder yore",  # No month or day, should return original string
        20250125,  # Not a valid date string, should return original integer
        None,  # None value, should return None
        "",  # Empty string, should return empty string
    ],
)
def test_should_never_parse(input_value: Any, relaxed_about_day: bool):
    """
    Tests for inputs which should never be parsed as datetimes.
    """
    result = gently_parse_datetime(input_value, relaxed_about_day=relaxed_about_day)
    assert result == input_value


@pytest.mark.parametrize(
    "input_value, expected_output",
    [
        ("October 2025", datetime(2025, 10, 1, tzinfo=timezone.utc)),
        ("2020 Apr", datetime(2020, 4, 1, tzinfo=timezone.utc)),
    ],
)
def test_should_parse_only_when_relaxed(input_value: Any, expected_output: datetime):
    """
    Tests for inputs which should only be parsed as datetimes when relaxed_about_day is True.
    """
    # Test with relaxed_about_day=True (must parse)
    result_relaxed = gently_parse_datetime(input_value, relaxed_about_day=True)
    assert result_relaxed == expected_output

    # Test with relaxed_about_day=False (must not parse)
    result_strict = gently_parse_datetime(input_value, relaxed_about_day=False)
    assert result_strict == input_value


# ---------------------------------------------------------------------------
# _cleanse_data
# ---------------------------------------------------------------------------


def test_cleanse_data_all_none_column_is_unchanged():
    data = {"col": [None, None, None]}
    assert _cleanse_data(data) == {"col": [None, None, None]}


def test_cleanse_data_uniform_type_is_unchanged():
    data = {"a": [1, 2, 3], "b": ["x", "y", "z"]}
    assert _cleanse_data(data) == data


def test_cleanse_data_majority_float_with_minority_int_is_unchanged():
    # float is majority; source_type=(float,int); all values satisfy → no coercion
    data = {"col": [1.0, 2.0, 3]}
    assert _cleanse_data(data) == data


def test_cleanse_data_mixed_types_with_no_date_falls_back_to_string():
    # majority float, but 'three' doesn't satisfy (float, int) → stringify all
    data = {"col": [1.0, 2.0, "three"]}
    assert _cleanse_data(data) == snapshot({"col": ["1.0", "2.0", "three"]})


def test_cleanse_data_date_column_with_parseable_string_parse_dates_true():
    dt = datetime(2023, 1, 15, tzinfo=timezone.utc)
    # '2023-06-15': day=15 > 12 so unambiguous regardless of dayfirst
    data = {"col": [dt, None, "2023-06-15"]}
    result = _cleanse_data(
        data, DateParsingOptions(parse_dates=True, relaxed_about_day=False)
    )
    assert result["col"][0] == dt
    assert result["col"][1] is None
    assert result["col"][2] == datetime(2023, 6, 15, tzinfo=timezone.utc)


def test_cleanse_data_date_column_with_parseable_string_parse_dates_false_falls_back_to_string():
    dt = datetime(2023, 1, 15, tzinfo=timezone.utc)
    data = {"col": [dt, "2023-06-15"]}
    result = _cleanse_data(data, DateParsingOptions(parse_dates=False))
    # parse_dates=False skips the date branch; mixed datetime+str → all stringified
    assert all(isinstance(v, str) for v in result["col"])


def test_cleanse_data_date_column_with_nones_are_preserved():
    dt = datetime(2023, 3, 1, tzinfo=timezone.utc)
    data = {"col": [dt, None, "2023-06-15"]}
    result = _cleanse_data(data, DateParsingOptions(parse_dates=True))
    assert result["col"][1] is None


def test_cleanse_data_handles_multiple_columns_independently():
    dt = datetime(2023, 1, 1, tzinfo=timezone.utc)
    data = {
        "dates": [dt, "2023-06-15"],  # day 15 is unambiguous
        "mixed": [1.0, "oops"],
    }
    result = _cleanse_data(data, DateParsingOptions(parse_dates=True))
    assert result["dates"][1] == datetime(2023, 6, 15, tzinfo=timezone.utc)
    assert result["mixed"] == ["1.0", "oops"]


# ---------------------------------------------------------------------------
# _dataframe_from_range
# ---------------------------------------------------------------------------


def _make_ws(headers: list[str], rows: list[list[Any]]):
    """Build a minimal in-memory openpyxl worksheet."""
    wb = OpxlWorkbook()
    ws = wb.active
    assert ws is not None
    for col_i, header in enumerate(headers, 1):
        ws.cell(1, col_i, header)
    for row_i, row in enumerate(rows, 2):
        for col_i, val in enumerate(row, 1):
            ws.cell(row_i, col_i, val)
    return ws


def test_dataframe_from_range_basic():
    ws = _make_ws(["Name", "Score"], [["Alice", 90], ["Bob", 85]])
    df = _dataframe_from_range(ws, "A1:B3")
    assert list(df.columns) == snapshot(["Name", "Score"])
    assert df["Name"].to_list() == snapshot(["Alice", "Bob"])
    assert df["Score"].to_list() == snapshot([90, 85])


def test_dataframe_from_range_normalises_newlines_in_column_headers():
    ws = _make_ws(["Full\nName", "Date\r\nIssued"], [["Alice", "2024-01-01"]])
    df = _dataframe_from_range(ws, "A1:B2")
    assert list(df.columns) == snapshot(["Full Name", "Date Issued"])


def test_dataframe_from_range_raises_type_error_for_merged_header_cell():
    wb = OpxlWorkbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "Name"
    ws["B1"] = "Value"
    ws["A2"] = "Alice"
    ws["B2"] = 42
    ws.merge_cells("A1:B1")  # B1 becomes a MergedCell
    with pytest.raises(TypeError):
        _dataframe_from_range(ws, "A1:B2")


def test_dataframe_from_range_raises_type_error_for_merged_data_cell():
    wb = OpxlWorkbook()
    ws = wb.active
    assert ws is not None
    ws["A1"] = "Name"
    ws["B1"] = "Value"
    ws["A2"] = "Alice"
    ws["B2"] = 42
    ws.merge_cells("A2:B2")  # B2 becomes a MergedCell
    with pytest.raises(TypeError):
        _dataframe_from_range(ws, "A1:B2")
